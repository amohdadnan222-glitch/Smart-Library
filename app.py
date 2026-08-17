from flask import Flask, render_template, request, redirect, session, flash
from datetime import datetime, timedelta
import sqlite3
from flask import send_file

from reportlab.platypus import SimpleDocTemplate
from reportlab.platypus import Table
from reportlab.platypus import TableStyle

from reportlab.lib import colors

from openpyxl import Workbook

import os
from werkzeug.security import generate_password_hash, check_password_hash

import re

app = Flask(__name__) 
app.secret_key = "library_management_secret"

@app.route("/")
def home():
    return render_template("index.html")

@app.route("/student_login", methods=["GET", "POST"])
def student_login():

    if request.method == "POST":

        email = request.form["email"]
        password = request.form["password"]

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        cursor.execute(
            "SELECT * FROM students WHERE email=?",
            (email,)
        )

        student = cursor.fetchone()

        conn.close()

        if student and check_password_hash(student[3], password):

            session["student_id"] = student[0]
            session["student_name"] = student[1]

            flash("Login Successful!", "success")

            return redirect("/student_dashboard")

        else:

            flash("Invalid Email or Password!", "danger")

            return redirect("/student_login")

    # This handles GET requests
    return render_template("student_login.html")


        


@app.route("/student_register", methods=["GET", "POST"])
def student_register():

    if request.method == "POST":

        name = request.form["name"]
        email = request.form["email"].strip().lower()
        password = request.form["password"]


        if not re.fullmatch(r"[A-Za-z0-9._%+-]+@gmail\.com", email):
            flash("Please enter a valid email address.", "danger")
            return redirect("/student_register")

        if len(password) < 6:
            flash("Password must be at least 6 characters long.", "danger")
            return redirect("/student_register")
        


        confirm_password = request.form["confirm_password"]
        hashed_password = generate_password_hash(password)

        if password != confirm_password:

            flash("Passwords do not match!", "danger")

            return redirect("/student_register")

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        # Check whether email already exists
        cursor.execute("""
            SELECT *
            FROM students
            WHERE email=?
        """, (email,))

        student = cursor.fetchone()

        if student:

            conn.close()

            flash("Email already registered!", "warning")

            return redirect("/student_register")

        cursor.execute("""
            INSERT INTO students
            (name,email,password)
            VALUES (?,?,?)
        """, (
            name,
            email,
            hashed_password
        ))

        conn.commit()

        conn.close()

        flash("Registration Successful! Please Login.", "success")

        return redirect("/student_login")

    return render_template("student_register.html")


@app.route("/student_dashboard")
def student_dashboard():

    if "student_id" not in session:
        return redirect("/student_login")

    return render_template(
        "student_dashboard.html",
        student_name=session["student_name"]
    )


@app.route("/student_profile")
def student_profile():

    if "student_id" not in session:
        return redirect("/student_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id, name, email
        FROM students
        WHERE id=?
    """, (session["student_id"],))

    student = cursor.fetchone()

    conn.close()

    return render_template(
        "student_profile.html",
        student=student
    )


@app.route("/edit_student_profile", methods=["GET", "POST"])
def edit_student_profile():

    if "student_id" not in session:
        return redirect("/student_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        name = request.form["name"]

        cursor.execute("""
            UPDATE students
            SET name=?
            WHERE id=?
        """, (
            name,
            session["student_id"]
        ))

        conn.commit()

        conn.close()

        flash("Profile updated successfully!", "success")

        return redirect("/student_profile")

    cursor.execute("""
        SELECT name,email
        FROM students
        WHERE id=?
    """, (session["student_id"],))

    student = cursor.fetchone()

    conn.close()

    return render_template(
        "edit_student_profile.html",
        student=student
    )



@app.route("/change_student_password", methods=["GET", "POST"])
def change_student_password():

    if "student_id" not in session:
        return redirect("/student_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        old_password = request.form["old_password"]
        new_password = request.form["new_password"]
        confirm_password = request.form["confirm_password"]

        cursor.execute(
            """
            SELECT password
            FROM students
            WHERE id=?
            """,
            (session["student_id"],)
        )

        student = cursor.fetchone()

        if not check_password_hash(student[0], old_password):

            conn.close()

            flash("Old password is incorrect!", "danger")

            return redirect("/change_student_password")

        if new_password != confirm_password:

            conn.close()

            flash("New passwords do not match!", "danger")

            return redirect("/change_student_password")

        hashed_password = generate_password_hash(new_password)

        cursor.execute(
            """
            UPDATE students
            SET password=?
            WHERE id=?
            """,
            (
                hashed_password,
                session["student_id"]
            )
        )

        conn.commit()

        conn.close()

        flash("Password changed successfully!", "success")

        return redirect("/student_profile")

    conn.close()

    return render_template("change_student_password.html")





@app.route("/teacher_login", methods=["GET", "POST"])
def teacher_login():

    if request.method == "POST":

        email = request.form["email"]
        password = request.form["password"]

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        cursor.execute(
        "SELECT * FROM teachers WHERE email=?",
        (email,)
        )

        teacher = cursor.fetchone()

        if teacher and check_password_hash(teacher[3], password):

            session["teacher_id"] = teacher[0]
            session["teacher_name"] = teacher[1]

            flash("Login Successful!", "success")

            conn.close()

            return redirect("/teacher_dashboard")

        else:

            conn.close()

            flash("Invalid Email or Password!", "danger")

            return redirect("/teacher_login")
        
    return render_template("teacher_login.html")

        


@app.route("/teacher_register", methods=["GET", "POST"])
def teacher_register():

    if request.method == "POST":

        name = request.form["name"]
        email = request.form["email"].strip()
        password = request.form["password"]


        if not re.fullmatch(r"[A-Za-z0-9._%+-]+@gmail\.com", email):
            flash("Please enter a valid email address.", "danger")
            return redirect("/teacher_register")

        if len(password) < 6:
            flash("Password must be at least 6 characters long.", "danger")
            return redirect("/teacher_register")
        
        confirm_password = request.form["confirm_password"]
        hashed_password = generate_password_hash(password)

        if password != confirm_password:

            flash("Passwords do not match!", "danger")
            return redirect("/teacher_register")

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        # Check duplicate email
        cursor.execute("""
            SELECT *
            FROM teachers
            WHERE email=?
        """, (email,))

        teacher = cursor.fetchone()

        if teacher:

            conn.close()

            flash("Email already registered!", "warning")

            return redirect("/teacher_register")

        cursor.execute("""
            INSERT INTO teachers
            (name,email,password)
            VALUES (?,?,?)
        """, (
            name,
            email,
            hashed_password
        ))

        conn.commit()
        conn.close()

        flash("Registration Successful! Please Login.", "success")

        return redirect("/teacher_login")

    return render_template("teacher_register.html")



@app.route("/teacher_dashboard")
def teacher_dashboard():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    return render_template(
        "teacher_dashboard.html",
        teacher_name=session["teacher_name"]
    )



@app.route("/teacher_profile")
def teacher_profile():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id, name, email
        FROM teachers
        WHERE id=?
    """, (session["teacher_id"],))

    teacher = cursor.fetchone()

    conn.close()

    return render_template(
        "teacher_profile.html",
        teacher=teacher
    )


@app.route("/edit_teacher_profile", methods=["GET", "POST"])
def edit_teacher_profile():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        name = request.form["name"]

        cursor.execute("""
            UPDATE teachers
            SET name=?
            WHERE id=?
        """, (
            name,
            session["teacher_id"]
        ))

        conn.commit()

        conn.close()

        flash("Profile updated successfully!", "success")

        return redirect("/teacher_profile")

    cursor.execute("""
        SELECT name,email
        FROM teachers
        WHERE id=?
    """, (session["teacher_id"],))

    teacher = cursor.fetchone()

    conn.close()

    return render_template(
        "edit_teacher_profile.html",
        teacher=teacher
    )


@app.route("/change_teacher_password", methods=["GET", "POST"])
def change_teacher_password():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        old_password = request.form["old_password"]
        new_password = request.form["new_password"]
        confirm_password = request.form["confirm_password"]

        cursor.execute("""
            SELECT password
            FROM teachers
            WHERE id=?
        """, (session["teacher_id"],))

        teacher = cursor.fetchone()

        if not check_password_hash(teacher[0], old_password):

            conn.close()

            flash("Old password is incorrect!", "danger")

            return redirect("/change_teacher_password")

        if new_password != confirm_password:

            conn.close()

            flash("New passwords do not match!", "danger")

            return redirect("/change_teacher_password")

        hashed_password = generate_password_hash(new_password)

        cursor.execute("""
            UPDATE teachers
            SET password=?
            WHERE id=?
        """, (
            hashed_password,
            session["teacher_id"]
        ))

        conn.commit()

        conn.close()

        flash("Password changed successfully!", "success")

        return redirect("/teacher_profile")

    conn.close()

    return render_template("change_teacher_password.html")


@app.route("/teacher_books")
def teacher_books():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    search = request.args.get("search", "")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT *
        FROM books
        WHERE
            title LIKE ?
            OR author LIKE ?
            OR isbn LIKE ?
            OR category LIKE ?
    """, (

        "%" + search + "%",
        "%" + search + "%",
        "%" + search + "%",
        "%" + search + "%"

    ))

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "teacher_books.html",
        books=books,
        search=search
    )




@app.route("/teacher_issue_book/<int:book_id>")
def teacher_issue_book(book_id):

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute(
        "SELECT available_quantity FROM books WHERE id=?",
        (book_id,)
    )

    book = cursor.fetchone()

    if book is None:
        conn.close()
        flash("Book not found!", "danger")
        return redirect("/teacher_books")

    if book[0] <= 0:
        conn.close()
        flash("Book not available!", "warning")
        return redirect("/teacher_books")

    issue_date = datetime.now().strftime("%Y-%m-%d")

    due_date = (datetime.now() + timedelta(days=15)).strftime("%Y-%m-%d")

    cursor.execute("""
        INSERT INTO issued_books
        (
            user_type,
            user_id,
            book_id,
            issue_date,
            due_date,
            status
        )
        VALUES
        (
            ?, ?, ?, ?, ?, ?
        )
    """,
    (
        "Teacher",
        session["teacher_id"],
        book_id,
        issue_date,
        due_date,
        "Issued"
    ))

    cursor.execute("""
        UPDATE books
        SET available_quantity=available_quantity-1
        WHERE id=?
    """,
    (book_id,))

    conn.commit()
    conn.close()

    flash("Book Issued Successfully!", "success")

    return redirect("/teacher_books")



@app.route("/teacher_search_books", methods=["GET", "POST"])
def teacher_search_books():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    books = []
    searched = False

    if request.method == "POST":

        keyword = request.form["keyword"]

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        cursor.execute("""
            SELECT *
            FROM books
            WHERE
                title LIKE ?
                OR author LIKE ?
                OR category LIKE ?
        """,
        (
            "%" + keyword + "%",
            "%" + keyword + "%",
            "%" + keyword + "%"
        ))

        books = cursor.fetchall()

        conn.close()

        searched = True

    return render_template(
        "teacher_search_books.html",
        books=books,
        searched=searched
    )




@app.route("/teacher_my_books")
def teacher_my_books():

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            books.title,
            issued_books.issue_date,
            issued_books.due_date,
            issued_books.return_date,
            issued_books.status,
            issued_books.fine,
            issued_books.id

        FROM issued_books

        JOIN books
        ON books.id = issued_books.book_id

        WHERE
            issued_books.user_id = ?
            AND issued_books.user_type = 'Teacher'

        ORDER BY issued_books.issue_date DESC
    """, (session["teacher_id"],))

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "teacher_my_books.html",
        books=books
    )





@app.route("/teacher_return_book/<int:issue_id>")
def teacher_return_book(issue_id):

    if "teacher_id" not in session:
        return redirect("/teacher_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            book_id,
            due_date
        FROM issued_books
        WHERE id=?
    """, (issue_id,))

    data = cursor.fetchone()

    if data is None:
        conn.close()
        flash("Record not found!", "danger")
        return redirect("/teacher_my_books")

    book_id = data[0]
    due_date = data[1]

    # Get today's date
    today = datetime.now().date()

    # Convert due date from text to date
    due = datetime.strptime(due_date, "%Y-%m-%d").date()

    # Calculate late days
    late_days = (today - due).days

    # Calculate fine
    fine = 0

    if late_days > 0:
        fine = late_days * 5

    # Update issued book
    cursor.execute("""
        UPDATE issued_books
        SET
            return_date=?,
            status=?,
            fine=?
        WHERE id=?
    """,
    (
        today.strftime("%Y-%m-%d"),
        "Returned",
        fine,
        issue_id
    ))

    # Increase available book quantity
    cursor.execute("""
        UPDATE books
        SET available_quantity=available_quantity+1
        WHERE id=?
    """, (book_id,))

    conn.commit()
    conn.close()

    flash(
        f"Book Returned Successfully! Fine = ₹{fine}",
        "success"
    )

    return redirect("/teacher_my_books")




@app.route("/admin_login", methods=["GET", "POST"])
def admin_login():

    if request.method == "POST":

        username = request.form["username"]
        password = request.form["password"]

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT *
            FROM admin
            WHERE username=?
            """,
            (username,)
        )

        admin = cursor.fetchone()

        conn.close()

        if admin and check_password_hash(admin[2], password):

            session["admin"] = admin[0]
            session["admin_username"] = admin[1]

            flash("Welcome Admin!", "success")

            return redirect("/admin_dashboard")

        flash("Invalid Username or Password!", "danger")

        return redirect("/admin_login")

    return render_template("admin_login.html")



@app.route("/admin_dashboard")
def admin_dashboard():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    # Dashboard Cards
    cursor.execute("SELECT COUNT(*) FROM books")
    total_books = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM students")
    total_students = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM teachers")
    total_teachers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM issued_books WHERE status='Issued'")
    total_issued = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM issued_books WHERE status='Returned'")
    total_returned = cursor.fetchone()[0]

    cursor.execute("""
        SELECT COUNT(*)
        FROM issued_books
        WHERE status='Issued'
        AND due_date < DATE('now')
    """)
    total_overdue = cursor.fetchone()[0]

    cursor.execute("SELECT IFNULL(SUM(fine),0) FROM issued_books")
    total_fine = cursor.fetchone()[0]

    # Books by Category
    cursor.execute("""
        SELECT category, COUNT(*)
        FROM books
        GROUP BY category
    """)
    category_data = cursor.fetchall()

    categories = [row[0] for row in category_data]
    category_counts = [row[1] for row in category_data]

    # Monthly Issues
    cursor.execute("""
        SELECT
            strftime('%m', issue_date),
            COUNT(*)
        FROM issued_books
        GROUP BY strftime('%m', issue_date)
        ORDER BY strftime('%m', issue_date)
    """)
    monthly_data = cursor.fetchall()

    months = [row[0] for row in monthly_data]
    monthly_count = [row[1] for row in monthly_data]

    conn.close()

    print("Categories:", categories)
    print("Category Counts:", category_counts)

    print("Months:", months)
    print("Monthly Count:", monthly_count)

    print("Issued:", total_issued)
    print("Returned:", total_returned)

    return render_template(
        "admin_dashboard.html",

        total_books=total_books,
        total_students=total_students,
        total_teachers=total_teachers,
        total_issued=total_issued,
        total_returned=total_returned,
        total_overdue=total_overdue,
        total_fine=total_fine,

        categories=categories,
        category_counts=category_counts,

        months=months,
        monthly_count=monthly_count
    )


@app.route("/admin_reports")
def admin_reports():

    if "admin" not in session:
        return redirect("/admin_login")

    return render_template("admin_reports.html")





@app.route("/add_book", methods=["GET", "POST"])
def add_book():

    if request.method == "POST":

        title = request.form["title"]
        author = request.form["author"]
        isbn = request.form["isbn"]
        category = request.form["category"]
        quantity = int(request.form["quantity"])

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        cursor.execute("""
        INSERT INTO books
        (title,author,isbn,category,quantity,available_quantity)
        VALUES (?,?,?,?,?,?)
        """,
        (title, author, isbn, category, quantity, quantity))

        conn.commit()
        conn.close()

        flash("Book added successfully!", "success")

        return redirect("/view_books")

    # GET request
    return render_template("add_book.html")



@app.route("/view_books")
def view_books():

    conn = sqlite3.connect("library.db")

    cursor = conn.cursor()

    cursor.execute("SELECT * FROM books")

    books = cursor.fetchall()

    conn.close()

    return render_template("view_books.html", books=books)





@app.route("/edit_book/<int:id>", methods=["GET", "POST"])
def edit_book(id):

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        title = request.form["title"]
        author = request.form["author"]
        isbn = request.form["isbn"]
        category = request.form["category"]
        quantity = int(request.form["quantity"])

        cursor.execute("""
        UPDATE books
        SET title=?,
            author=?,
            isbn=?,
            category=?,
            quantity=?,
            available_quantity=?
        WHERE id=?
        """,
        (title,
         author,
         isbn,
         category,
         quantity,
         quantity,
         id))

        conn.commit()
        conn.close()

        flash("Book updated successfully!", "info")
        return redirect("/view_books")

    cursor.execute("SELECT * FROM books WHERE id=?", (id,))
    book = cursor.fetchone()

    conn.close()

    return render_template("edit_book.html", book=book)





@app.route("/delete_book/<int:id>")
def delete_book(id):

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("DELETE FROM books WHERE id=?", (id,))

    conn.commit()
    conn.close()

    flash("Book deleted successfully!", "warning")
    return redirect("/view_books")





@app.route("/student_books")
def student_books():

    if "student_id" not in session:
        return redirect("/student_login")

    search = request.args.get("search", "")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT *
        FROM books
        WHERE
            title LIKE ?
            OR author LIKE ?
            OR isbn LIKE ?
            OR category LIKE ?
    """, (

        "%" + search + "%",
        "%" + search + "%",
        "%" + search + "%",
        "%" + search + "%"

    ))

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "student_books.html",
        books=books,
        search=search
    )





@app.route("/my_books")
def my_books():

    if "student_id" not in session:
        return redirect("/student_login")

    student_id = session["student_id"]

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            books.title,
            issued_books.issue_date,
            issued_books.return_date,
            issued_books.due_date,
            issued_books.status,
            issued_books.fine,
            issued_books.id

        FROM issued_books

        JOIN books
        ON books.id = issued_books.book_id

        WHERE issued_books.user_id = ?
        AND issued_books.user_type = 'Student'
    """, (student_id,))

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "my_books.html",
        books=books
    )





@app.route("/return_book/<int:issue_id>")
def return_book(issue_id):

    if "student_id" not in session:
        return redirect("/student_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            book_id,
            due_date
        FROM issued_books
        WHERE id=?
    """,(issue_id,))

    data = cursor.fetchone()

    if data is None:
        conn.close()
        flash("Record not found!", "danger")
        return redirect("/my_books")

    book_id = data[0]
    due_date = data[1]

    today = datetime.now().date()
    due = datetime.strptime(due_date,"%Y-%m-%d").date()

    late_days = (today-due).days

    fine = 0

    if late_days > 0:
        fine = late_days * 5

    cursor.execute("""
        UPDATE issued_books
        SET
            status=?,
            fine=?
        WHERE id=?
    """,
    (
        "Returned",
        fine,
        issue_id
    ))

    cursor.execute("""
        UPDATE books
        SET available_quantity=available_quantity+1
        WHERE id=?
    """,(book_id,))

    conn.commit()
    conn.close()

    flash(f"Book Returned Successfully! Fine = ₹{fine}", "success")

    return redirect("/my_books")





@app.route("/issued_books")
def issued_books():

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
    SELECT
        issued_books.id,
        CASE
        WHEN issued_books.user_type='Student'
        THEN students.name
        ELSE teachers.name
        END AS user_name,

    books.title,

    issued_books.user_type,

    issued_books.issue_date,

    issued_books.return_date,

    issued_books.status

    FROM issued_books

        JOIN books
        ON books.id = issued_books.book_id

    LEFT JOIN students
    ON students.id = issued_books.user_id
    AND issued_books.user_type='Student'

    LEFT JOIN teachers
    ON teachers.id = issued_books.user_id
    AND issued_books.user_type='Teacher'
        """)

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "issued_books.html",
        books=books
    )




@app.route("/issue_book/<int:book_id>")
def issue_book(book_id):

    if "student_id" not in session:
        return redirect("/student_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute(
        "SELECT available_quantity FROM books WHERE id=?",
        (book_id,)
    )

    book = cursor.fetchone()

    if book is None:
        conn.close()
        flash("Book not found!", "danger")
        return redirect("/student_books")

    if book[0] <= 0:
        conn.close()
        flash("Book not available!", "warning")
        return redirect("/student_books")

    # Issue Date
    issue_date = datetime.now().strftime("%Y-%m-%d")

    # Due Date (15 Days)
    due_date = (datetime.now() + timedelta(days=15)).strftime("%Y-%m-%d")

    cursor.execute("""
        INSERT INTO issued_books
        (
            user_type,
            user_id,
            book_id,
            issue_date,
            due_date,
            status
        )
        VALUES
        (
            ?, ?, ?, ?, ?, ?
        )
    """,
    (
        "Student",
        session["student_id"],
        book_id,
        issue_date,
        due_date,
        "Issued"
    ))

    cursor.execute("""
        UPDATE books
        SET available_quantity=available_quantity-1
        WHERE id=?
    """,
    (book_id,))

    conn.commit()
    conn.close()

    flash("Book Issued Successfully!", "success")

    return redirect("/student_books")





@app.route("/search_books", methods=["GET", "POST"])
def search_books():

    if "student_id" not in session:
        return redirect("/student_login")

    books = []
    searched = False

    if request.method == "POST":

        keyword = request.form["keyword"]

        conn = sqlite3.connect("library.db")
        cursor = conn.cursor()

        cursor.execute("""

        SELECT *

        FROM books

        WHERE

        title LIKE ?

        OR author LIKE ?

        OR category LIKE ?

        """,

        (
            "%" + keyword + "%",
            "%" + keyword + "%",
            "%" + keyword + "%"
        ))

        books = cursor.fetchall()

        conn.close()

        searched = True

    return render_template(
        "search_books.html",
        books=books,
        searched=searched
    )




@app.route("/admin_issue_history")
def admin_issue_history():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get("search", "")
    status = request.args.get("status", "")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    query = """
    SELECT

        issued_books.id,

        CASE
            WHEN issued_books.user_type='Student'
            THEN students.name
            ELSE teachers.name
        END,

        issued_books.user_type,

        books.title,

        issued_books.issue_date,

        issued_books.due_date,

        issued_books.return_date,

        issued_books.status,

        issued_books.fine

    FROM issued_books

    JOIN books
    ON books.id=issued_books.book_id

    LEFT JOIN students
    ON students.id=issued_books.user_id
    AND issued_books.user_type='Student'

    LEFT JOIN teachers
    ON teachers.id=issued_books.user_id
    AND issued_books.user_type='Teacher'

    WHERE
    (
        books.title LIKE ?
        OR students.name LIKE ?
        OR teachers.name LIKE ?
        OR issued_books.user_type LIKE ?
    )
    """

    parameters = [
        "%" + search + "%",
        "%" + search + "%",
        "%" + search + "%",
        "%" + search + "%"
    ]

    if status != "":
        query += " AND issued_books.status=?"
        parameters.append(status)

    query += " ORDER BY issued_books.id DESC"

    cursor.execute(query, parameters)

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "admin_issue_history.html",
        books=books,
        search=search,
        status=status
    )


@app.route("/manage_students")
def manage_students():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            name,
            email
        FROM students
        ORDER BY name
    """)

    students = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_students.html",
        students=students
    )


@app.route("/edit_student/<int:id>", methods=["GET", "POST"])
def edit_student(id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        name = request.form["name"]
        email = request.form["email"]
        password = request.form["password"]

        cursor.execute("""
            UPDATE students
            SET
                name=?,
                email=?,
                password=?
            WHERE id=?
        """, (
            name,
            email,
            password,
            id
        ))

        conn.commit()
        conn.close()

        flash("Student updated successfully!", "success")

        return redirect("/manage_students")

    cursor.execute("""
        SELECT *
        FROM students
        WHERE id=?
    """, (id,))

    student = cursor.fetchone()

    conn.close()

    return render_template(
        "edit_student.html",
        student=student
    )


@app.route("/delete_student/<int:id>")
def delete_student(id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    # Check if student exists
    cursor.execute("""
        SELECT *
        FROM students
        WHERE id=?
    """, (id,))

    student = cursor.fetchone()

    if student is None:

        conn.close()

        flash("Student not found!", "danger")

        return redirect("/manage_students")

    # Delete student
    cursor.execute("""
        DELETE FROM students
        WHERE id=?
    """, (id,))

    conn.commit()

    conn.close()

    flash("Student deleted successfully!", "success")

    return redirect("/manage_students")


@app.route("/student_history/<int:id>")
def student_history(id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    # Get Student Name
    cursor.execute("""
        SELECT name
        FROM students
        WHERE id=?
    """, (id,))

    student = cursor.fetchone()

    if student is None:

        conn.close()

        flash("Student not found!", "danger")

        return redirect("/manage_students")

    student_name = student[0]

    # Get Issue History
    cursor.execute("""
        SELECT
            books.title,
            issued_books.issue_date,
            issued_books.due_date,
            issued_books.return_date,
            issued_books.status,
            issued_books.fine

        FROM issued_books

        JOIN books
        ON books.id = issued_books.book_id

        WHERE
            issued_books.user_id=?
            AND issued_books.user_type='Student'

        ORDER BY issued_books.issue_date DESC
    """, (id,))

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "student_history.html",
        books=books,
        student_name=student_name
    )


@app.route("/manage_teachers")
def manage_teachers():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            name,
            email
        FROM teachers
        ORDER BY name
    """)

    teachers = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_teachers.html",
        teachers=teachers
    )


@app.route("/edit_teacher/<int:id>", methods=["GET", "POST"])
def edit_teacher(id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    if request.method == "POST":

        name = request.form["name"]
        email = request.form["email"]
        password = request.form["password"]

        cursor.execute("""
            UPDATE teachers
            SET
                name=?,
                email=?,
                password=?
            WHERE id=?
        """, (
            name,
            email,
            password,
            id
        ))

        conn.commit()
        conn.close()

        flash("Teacher updated successfully!", "success")

        return redirect("/manage_teachers")

    cursor.execute("""
        SELECT *
        FROM teachers
        WHERE id=?
    """, (id,))

    teacher = cursor.fetchone()

    conn.close()

    return render_template(
        "edit_teacher.html",
        teacher=teacher
    )


@app.route("/delete_teacher/<int:id>")
def delete_teacher(id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    # Check whether teacher exists
    cursor.execute("""
        SELECT *
        FROM teachers
        WHERE id=?
    """, (id,))

    teacher = cursor.fetchone()

    if teacher is None:

        conn.close()

        flash("Teacher not found!", "danger")

        return redirect("/manage_teachers")

    # Delete Teacher
    cursor.execute("""
        DELETE FROM teachers
        WHERE id=?
    """, (id,))

    conn.commit()

    conn.close()

    flash("Teacher deleted successfully!", "success")

    return redirect("/manage_teachers")


@app.route("/teacher_history/<int:id>")
def teacher_history(id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    # Get teacher name
    cursor.execute("""
        SELECT name
        FROM teachers
        WHERE id=?
    """, (id,))

    teacher = cursor.fetchone()

    if teacher is None:

        conn.close()

        flash("Teacher not found!", "danger")

        return redirect("/manage_teachers")

    teacher_name = teacher[0]

    # Get issue history
    cursor.execute("""
        SELECT
            books.title,
            issued_books.issue_date,
            issued_books.due_date,
            issued_books.return_date,
            issued_books.status,
            issued_books.fine

        FROM issued_books

        JOIN books
        ON books.id = issued_books.book_id

        WHERE
            issued_books.user_id=?
            AND issued_books.user_type='Teacher'

        ORDER BY issued_books.issue_date DESC
    """, (id,))

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "teacher_history.html",
        books=books,
        teacher_name=teacher_name
    )


@app.route("/download_issue_report")
def download_issue_report():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""

        SELECT

            books.title,

            issued_books.user_type,

            issued_books.issue_date,

            issued_books.due_date,

            issued_books.return_date,

            issued_books.status,

            issued_books.fine

        FROM issued_books

        JOIN books

        ON books.id=issued_books.book_id

    """)

    rows = cursor.fetchall()

    conn.close()

    filename = "Issued_Books_Report.pdf"

    doc = SimpleDocTemplate(filename)

    data = [[

        "Book",

        "User",

        "Issue",

        "Due",

        "Return",

        "Status",

        "Fine"

    ]]

    for row in rows:

        data.append(list(row))

    table = Table(data)

    table.setStyle(TableStyle([

        ("BACKGROUND",(0,0),(-1,0),colors.grey),

        ("TEXTCOLOR",(0,0),(-1,0),colors.white),

        ("GRID",(0,0),(-1,-1),1,colors.black),

        ("BACKGROUND",(0,1),(-1,-1),colors.beige)

    ]))

    doc.build([table])

    return send_file(

        filename,

        as_attachment=True

    )



@app.route("/export_books_excel")
def export_books_excel():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            title,
            author,
            isbn,
            category,
            quantity,
            available_quantity
        FROM books
        ORDER BY title
    """)

    books = cursor.fetchall()

    conn.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Books"

    # Header Row
    sheet.append([
        "Book ID",
        "Title",
        "Author",
        "ISBN",
        "Category",
        "Total Quantity",
        "Available Quantity"
    ])

    # Data Rows
    for book in books:
        sheet.append(book)

    filename = "Books.xlsx"

    workbook.save(filename)

    return send_file(
        filename,
        as_attachment=True
    )


@app.route("/export_students_excel")
def export_students_excel():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            name,
            email
        FROM students
        ORDER BY name
    """)

    students = cursor.fetchall()

    conn.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Students"

    sheet.append([
        "Student ID",
        "Name",
        "Email"
    ])

    for student in students:
        sheet.append(student)

    filename = "Students.xlsx"

    workbook.save(filename)

    return send_file(
        filename,
        as_attachment=True
    )


@app.route("/export_teachers_excel")
def export_teachers_excel():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            name,
            email
        FROM teachers
        ORDER BY name
    """)

    teachers = cursor.fetchall()

    conn.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Teachers"

    sheet.append([
        "Teacher ID",
        "Name",
        "Email"
    ])

    for teacher in teachers:
        sheet.append(teacher)

    filename = "Teachers.xlsx"

    workbook.save(filename)

    return send_file(
        filename,
        as_attachment=True
    )


@app.route("/report/books")
def books_report():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            title,
            author,
            category,
            quantity,
            available_quantity
        FROM books
        ORDER BY title
    """)

    books = cursor.fetchall()

    conn.close()

    return render_template(
        "books_report.html",
        books=books
    )



@app.route("/report/students")
def students_report():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            name,
            email
        FROM students
        ORDER BY name
    """)

    students = cursor.fetchall()

    conn.close()

    return render_template(
        "students_report.html",
        students=students
    )


@app.route("/report/teachers")
def teachers_report():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            name,
            email
        FROM teachers
        ORDER BY name
    """)

    teachers = cursor.fetchall()

    conn.close()

    return render_template(
        "teachers_report.html",
        teachers=teachers
    )



@app.route("/report/issues")
def issues_report():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("library.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT

            books.title,

            issued_books.user_type,

            issued_books.issue_date,

            issued_books.due_date,

            issued_books.return_date,

            issued_books.status,

            issued_books.fine

        FROM issued_books

        JOIN books

        ON books.id = issued_books.book_id

        ORDER BY issued_books.issue_date DESC
    """)

    issues = cursor.fetchall()

    conn.close()

    return render_template(
        "issues_report.html",
        issues=issues
    )



@app.route("/logout")
def logout():

    session.clear()

    flash("Logged out successfully!", "success")

    return redirect("/")



@app.errorhandler(404)
def page_not_found(error):
    return render_template("404.html"), 404



if __name__ == "__main__":
    app.run(debug=True)